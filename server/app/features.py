"""Lembretes, tarefas, documentos, contas fixas, convites, bom dia e alertas de gastos.

Regra de segurança mantida: nada sai em nome do usuário sem confirmação. Convidar pessoas para um
evento (o Google manda e-mail para elas) vira uma ação pendente, igual ao rascunho de e-mail.
"""
import calendar as _cal
import hashlib
import hmac
import time
import uuid
from datetime import date, datetime, timedelta
from zoneinfo import ZoneInfo

from . import config, google_client, store

PRIORITIES = ["alta", "normal", "baixa"]

TOOLS = [
    {
        "name": "create_reminder",
        "description": "Cria um LEMBRETE (alerta no celular na hora marcada, via Google Calendar). Use para "
                       "'me lembra de...', 'me avisa...'. Pode repetir (recurrence).",
        "parameters": {
            "type": "object",
            "properties": {
                "text": {"type": "string", "description": "o que lembrar, curto. Ex. 'Pagar o IVA'"},
                "when": {"type": "string", "description": "ISO 8601 local, ex. 2026-10-12T09:00:00"},
                "recurrence": {"type": "string", "description": "opcional, RRULE. Ex. 'RRULE:FREQ=MONTHLY;BYMONTHDAY=5'"},
                "notes": {"type": "string"},
            },
            "required": ["text", "when"],
        },
    },
    {
        "name": "add_meet_link",
        "description": "Gera um link do Google Meet num evento existente (event_id de list_calendar_events).",
        "parameters": {"type": "object", "properties": {"event_id": {"type": "string"}}, "required": ["event_id"]},
    },
    {
        "name": "prepare_event_invite",
        "description": "Prepara o convite de pessoas (e-mails) para um evento existente. NÃO envia: o usuário "
                       "confirma no app, e só então o Google manda o convite. Peça os e-mails se não souber.",
        "parameters": {
            "type": "object",
            "properties": {
                "event_id": {"type": "string"},
                "emails": {"type": "array", "items": {"type": "string"}},
            },
            "required": ["event_id", "emails"],
        },
    },
    {
        "name": "add_task",
        "description": "Cria uma tarefa (coisa a fazer, sem horário fixo de compromisso). Prazo opcional. "
                       "Se o usuário quiser ser avisado, passe remind_at e o celular alerta.",
        "parameters": {
            "type": "object",
            "properties": {
                "title": {"type": "string"},
                "due": {"type": "string", "description": "prazo AAAA-MM-DD (opcional)"},
                "priority": {"type": "string", "enum": PRIORITIES},
                "notes": {"type": "string"},
                "remind_at": {"type": "string", "description": "ISO 8601 local para alertar (opcional)"},
            },
            "required": ["title"],
        },
    },
    {
        "name": "list_tasks",
        "description": "Lista tarefas. status: 'aberta' (padrão), 'feita' ou 'todas'.",
        "parameters": {"type": "object", "properties": {"status": {"type": "string", "enum": ["aberta", "feita", "todas"]}}},
    },
    {
        "name": "complete_task",
        "description": "Marca uma tarefa como feita (task_id de list_tasks).",
        "parameters": {"type": "object", "properties": {"task_id": {"type": "integer"}}, "required": ["task_id"]},
    },
    {
        "name": "save_document",
        "description": "Guarda a foto desta mensagem como DOCUMENTO (seguro, contrato, carta, MOT, apólice, "
                       "identidade, garantia...). Descreva bem para achar depois. Se tiver data de validade, "
                       "passe expires_on: o Fidus avisa antes de vencer.",
        "parameters": {
            "type": "object",
            "properties": {
                "title": {"type": "string", "description": "ex. 'Seguro do Land Rover - Admiral'"},
                "description": {"type": "string", "description": "o que é, números importantes, empresa, pessoa"},
                "tags": {"type": "string", "description": "palavras-chave separadas por vírgula"},
                "expires_on": {"type": "string", "description": "AAAA-MM-DD, se houver validade"},
                "remind_days_before": {"type": "integer", "description": "padrão 30"},
            },
            "required": ["title", "description"],
        },
    },
    {
        "name": "search_documents",
        "description": "Procura documentos guardados por descrição ('seguro do carro', 'contrato aluguel'). "
                       "Sem query, lista os mais recentes. O app mostra o botão para abrir.",
        "parameters": {"type": "object", "properties": {"query": {"type": "string"}}},
    },
    {
        "name": "add_bill",
        "description": "Cadastra uma CONTA FIXA mensal (aluguel, seguro, internet, assinatura). O Fidus avisa "
                       "alguns dias antes de vencer, todo mês.",
        "parameters": {
            "type": "object",
            "properties": {
                "name": {"type": "string"},
                "day_of_month": {"type": "integer", "description": "dia do vencimento, 1-31"},
                "amount": {"type": "number"},
                "currency": {"type": "string"},
                "business": {"type": "string", "description": "uma das empresas/carteiras do usuário"},
                "category": {"type": "string", "enum": config.CATEGORIES},
                "remind_days_before": {"type": "integer", "description": "padrão 2"},
            },
            "required": ["name", "day_of_month"],
        },
    },
    {
        "name": "list_bills",
        "description": "Lista as contas fixas cadastradas e quando vence cada uma.",
        "parameters": {"type": "object", "properties": {}},
    },
    {
        "name": "delete_bill",
        "description": "Remove uma conta fixa (bill_id de list_bills) e o aviso mensal dela.",
        "parameters": {"type": "object", "properties": {"bill_id": {"type": "integer"}}, "required": ["bill_id"]},
    },
    {
        "name": "weekly_review",
        "description": "Resumo da semana: o que o Fidus fez, tarefas concluídas, gastos da semana x semana anterior, "
                       "e o que vem na próxima semana. Use para 'como foi minha semana', 'resumo semanal'.",
        "parameters": {"type": "object", "properties": {}},
    },
    {
        "name": "export_for_accountant",
        "description": "Gera um pacote para o contador: planilha Excel dos gastos (com IVA) + fotos dos recibos, num "
                       "zip. Filtra por período e empresa. O app mostra o botão para abrir/baixar.",
        "parameters": {
            "type": "object",
            "properties": {
                "date_from": {"type": "string", "description": "AAAA-MM-DD"},
                "date_to": {"type": "string", "description": "AAAA-MM-DD"},
                "business": {"type": "string", "description": "uma das empresas/carteiras do usuário"},
            },
            "required": ["date_from", "date_to"],
        },
    },
    {
        "name": "list_subscriptions",
        "description": "Detecta cobranças recorrentes (assinaturas, mensalidades) nos gastos lançados e mostra valor "
                       "atual, mudanças de preço e total por mês.",
        "parameters": {"type": "object", "properties": {}},
    },
    {
        "name": "daily_briefing",
        "description": "Resumo do dia: agenda, tarefas (atrasadas e do dia), contas a vencer, documentos "
                       "vencendo, e-mails importantes não lidos e gastos do mês. Use para 'bom dia', "
                       "'o que eu tenho hoje', 'como está meu dia'.",
        "parameters": {"type": "object", "properties": {"day": {"type": "string", "description": "AAAA-MM-DD, padrão hoje"}}},
    },
]


def _tz():
    return ZoneInfo(store.user_tz())


def _now() -> datetime:
    return datetime.now(_tz())


def _cal_api():
    return google_client.calendar()


def _event_body(title, start: datetime, minutes: int, description=None, recurrence=None, popup_minutes=0):
    tz = store.user_tz()
    body = {
        "summary": title,
        "start": {"dateTime": start.replace(tzinfo=None).isoformat(), "timeZone": tz},
        "end": {"dateTime": (start + timedelta(minutes=minutes)).replace(tzinfo=None).isoformat(), "timeZone": tz},
        "reminders": {"useDefault": False, "overrides": [{"method": "popup", "minutes": popup_minutes}]},
    }
    if description:
        body["description"] = description
    if recurrence:
        body["recurrence"] = [recurrence if recurrence.startswith("RRULE:") else f"RRULE:{recurrence}"]
    return body


def _insert(body) -> str:
    return _cal_api().events().insert(calendarId="primary", body=body).execute()["id"]


def _delete_event_quiet(event_id):
    if not event_id:
        return
    try:
        _cal_api().events().delete(calendarId="primary", eventId=event_id).execute()
    except Exception:
        pass


def _parse_local(iso: str) -> datetime:
    return datetime.fromisoformat(iso).replace(tzinfo=None)


# ---------- Lembretes e agenda ----------
def create_reminder(text, when, recurrence=None, notes=None):
    start = _parse_local(when)
    desc = ("Lembrete do Fidus." + (f"\n{notes}" if notes else ""))
    eid = _insert(_event_body(f"⏰ {text}", start, 15, desc, recurrence, popup_minutes=0))
    return {"ok": True, "event_id": eid, "text": text, "when": when, "recurring": bool(recurrence)}


def add_meet_link(event_id):
    ev = _cal_api().events().patch(
        calendarId="primary", eventId=event_id, conferenceDataVersion=1,
        body={"conferenceData": {"createRequest": {"requestId": uuid.uuid4().hex,
                                                   "conferenceSolutionKey": {"type": "hangoutsMeet"}}}}).execute()
    link = ev.get("hangoutLink")
    for ep in (ev.get("conferenceData") or {}).get("entryPoints", []):
        if ep.get("entryPointType") == "video":
            link = ep.get("uri")
    return {"ok": True, "event_id": event_id, "title": ev.get("summary"), "meet_link": link}


def prepare_event_invite(event_id, emails):
    emails = [e.strip() for e in emails if e and "@" in e]
    if not emails:
        return {"error": "nenhum e-mail válido. Pergunte o e-mail das pessoas."}
    ev = _cal_api().events().get(calendarId="primary", eventId=event_id).execute()
    payload = {"event_id": event_id, "title": ev.get("summary"),
               "start": ev.get("start", {}).get("dateTime") or ev.get("start", {}).get("date"),
               "emails": emails}
    pid = store.create_pending("calendar_invite", payload)
    return {"ok": True, "pending_action_id": pid, "status": "aguardando confirmação do usuário", "invite": payload}


def send_confirmed_invite(payload: dict) -> dict:
    """Chamado SOMENTE pelo endpoint de confirmação do app."""
    cal = _cal_api()
    ev = cal.events().get(calendarId="primary", eventId=payload["event_id"]).execute()
    attendees = ev.get("attendees", [])
    have = {a.get("email", "").lower() for a in attendees}
    attendees += [{"email": e} for e in payload["emails"] if e.lower() not in have]
    cal.events().patch(calendarId="primary", eventId=payload["event_id"], sendUpdates="all",
                       body={"attendees": attendees}).execute()
    return {"ok": True}


# ---------- Tarefas ----------
def add_task(title, due=None, priority="normal", notes=None, remind_at=None):
    eid = None
    if remind_at:
        eid = create_reminder(title, remind_at, notes=notes)["event_id"]
    tid = store.insert("tasks", title=title, due=due, priority=priority or "normal", notes=notes, event_id=eid)
    return {"ok": True, "task_id": tid, "title": title, "due": due, "priority": priority or "normal",
            "reminder": bool(eid)}


def list_tasks(status="aberta"):
    if status == "todas":
        rows = store.select("SELECT * FROM tasks ORDER BY status, due IS NULL, due, id DESC LIMIT 100")
    else:
        rows = store.select("SELECT * FROM tasks WHERE status=? ORDER BY due IS NULL, due, "
                            "CASE priority WHEN 'alta' THEN 0 WHEN 'normal' THEN 1 ELSE 2 END, id DESC LIMIT 100",
                            (status,))
    today = _now().date().isoformat()
    for r in rows:
        r["overdue"] = bool(r["status"] == "aberta" and r["due"] and r["due"] < today)
    return {"tasks": [{k: r[k] for k in ("id", "title", "due", "priority", "status", "notes", "overdue")} for r in rows]}


def complete_task(task_id):
    t = store.get("tasks", int(task_id))
    if not t:
        return {"error": "tarefa não encontrada"}
    store.update("tasks", int(task_id), status="feita", done_at=store.now())
    return {"ok": True, "task_id": int(task_id), "title": t["title"]}


def reopen_task(task_id):
    store.update("tasks", int(task_id), status="aberta", done_at=None)


def delete_task(task_id):
    t = store.get("tasks", int(task_id))
    if t:
        _delete_event_quiet(t.get("event_id"))
        store.update("tasks", int(task_id), status="apagada")


# ---------- Documentos ----------
def save_document(title, description, tags=None, expires_on=None, remind_days_before=30, _path=None):
    if not _path:
        return {"error": "nenhuma foto nesta mensagem. Peça para o usuário mandar a foto do documento."}
    eid = None
    if expires_on:
        d = date.fromisoformat(expires_on)
        when = datetime.combine(max(d - timedelta(days=int(remind_days_before or 30)), _now().date() + timedelta(days=1)),
                                datetime.min.time()).replace(hour=9)
        eid = _insert(_event_body(f"📄 Vence em {d:%d/%m/%Y}: {title}", when, 15,
                                  "Aviso do Fidus: documento perto de vencer.", popup_minutes=0))
    did = store.insert("documents", title=title, description=description, tags=tags, path=_path,
                       media_type="image", expires_on=expires_on, event_id=eid)
    return {"ok": True, "document_id": did, "title": title, "expires_on": expires_on, "reminder": bool(eid)}


def search_documents(query=None):
    if query:
        words = [w for w in query.lower().split() if len(w) > 2] or [query.lower()]
        rows = store.select("SELECT * FROM documents WHERE deleted=0 ORDER BY id DESC LIMIT 300")
        scored = []
        for r in rows:
            hay = " ".join(str(r.get(k) or "") for k in ("title", "description", "tags")).lower()
            score = sum(1 for w in words if w in hay)
            if score:
                scored.append((score, r))
        rows = [r for _, r in sorted(scored, key=lambda x: -x[0])[:5]]
    else:
        rows = store.select("SELECT * FROM documents WHERE deleted=0 ORDER BY id DESC LIMIT 10")
    return {"documents": [{"document_id": r["id"], "title": r["title"], "description": (r["description"] or "")[:300],
                           "expires_on": r["expires_on"]} for r in rows]}


def delete_document(doc_id):
    d = store.get("documents", int(doc_id))
    if d:
        _delete_event_quiet(d.get("event_id"))
        store.update("documents", int(doc_id), deleted=1)


def _doc_sig(uid: str, doc_id: int, exp: int) -> str:
    return hmac.new(config.APP_TOKEN.encode(), f"doc:{uid}:{doc_id}:{exp}".encode(), hashlib.sha256).hexdigest()[:32]


def sign(doc_id: int, ttl: int = 3600) -> str:
    """Link temporário para o documento do cliente atual (o id do cliente entra na assinatura)."""
    uid = store.current()["id"]
    exp = int(time.time()) + ttl
    return (f"{config.PUBLIC_BASE_URL.rstrip('/')}/v1/documents/{doc_id}/file"
            f"?u={uid}&exp={exp}&sig={_doc_sig(uid, doc_id, exp)}")


def check_sig(doc_id: int, exp: int, sig: str, uid: str = store.OWNER_ID) -> bool:
    if exp < time.time():
        return False
    good = _doc_sig(uid, doc_id, exp)
    try:
        return hmac.compare_digest(good, sig)
    except TypeError:
        return False


# ---------- Contas fixas ----------
def _next_due(day: int, today: date | None = None) -> date:
    today = today or _now().date()
    y, m = today.year, today.month
    for _ in range(2):
        d = date(y, m, min(day, _cal.monthrange(y, m)[1]))
        if d >= today:
            return d
        y, m = (y + 1, 1) if m == 12 else (y, m + 1)
    return d


def add_bill(name, day_of_month, amount=None, currency=None, business=None, category=None, remind_days_before=2):
    day = int(day_of_month)
    if not 1 <= day <= 31:
        return {"error": "dia do mês inválido"}
    # dia 29-31 não existe em todo mês: o aviso fica no máximo no dia 28 para não pular meses
    rday = min(max(day - int(remind_days_before if remind_days_before is not None else 2), 1), 28)
    first = _next_due(rday)
    cur = (currency or store.default_currency()).upper()
    business = store.match_business(business) or business
    label = f"🔁 Vence dia {day}: {name}" + (f" ({amount:.2f} {cur})" if amount else "")
    eid = _insert(_event_body(label, datetime.combine(first, datetime.min.time()).replace(hour=9), 15,
                              "Conta fixa cadastrada no Fidus.", f"RRULE:FREQ=MONTHLY;BYMONTHDAY={rday}"))
    bid = store.insert("bills", name=name, amount=amount, currency=cur, business=business, category=category,
                       day_of_month=day, event_id=eid)
    return {"ok": True, "bill_id": bid, "name": name, "day_of_month": day, "amount": amount, "currency": cur,
            "next_due": _next_due(day).isoformat()}


def list_bills():
    rows = store.select("SELECT * FROM bills WHERE active=1 ORDER BY day_of_month")
    return {"bills": [{"bill_id": r["id"], "name": r["name"], "amount": r["amount"], "currency": r["currency"],
                       "business": r["business"], "day_of_month": r["day_of_month"],
                       "next_due": _next_due(r["day_of_month"]).isoformat()} for r in rows]}


def delete_bill(bill_id):
    b = store.get("bills", int(bill_id))
    if not b:
        return {"error": "conta não encontrada"}
    _delete_event_quiet(b.get("event_id"))
    store.update("bills", int(bill_id), active=0)
    return {"ok": True, "bill_id": int(bill_id), "name": b["name"]}


# ---------- Gastos: alertas e comparação ----------
def expense_alerts(eid: int, amount: float, currency: str, category: str, merchant: str | None, day: str) -> list[str]:
    alerts = []
    d = date.fromisoformat(day)
    dup = store.select(
        "SELECT * FROM expenses WHERE deleted=0 AND id<>? AND currency=? AND abs(amount-?)<0.005 "
        "AND date>=? AND date<=?", (eid, currency, amount, (d - timedelta(days=3)).isoformat(),
                                   (d + timedelta(days=3)).isoformat()))
    dup = [x for x in dup if not merchant or not x["merchant"] or x["merchant"].lower() == merchant.lower()]
    if dup:
        alerts.append(f"possível duplicado: já existe {amount:.2f} {currency} em {dup[0]['date']}"
                      f"{' no ' + dup[0]['merchant'] if dup[0]['merchant'] else ''} (id {dup[0]['id']})")
    # assinatura que ficou mais cara: mesmo estabelecimento ~1 mês antes com valor menor
    if merchant:
        prev_m = store.select(
            "SELECT * FROM expenses WHERE deleted=0 AND id<>? AND currency=? AND lower(merchant)=lower(?) "
            "AND date>=? AND date<=? ORDER BY date DESC LIMIT 1",
            (eid, currency, merchant, (d - timedelta(days=40)).isoformat(), (d - timedelta(days=20)).isoformat()))
        if prev_m and amount > prev_m[0]["amount"] * 1.03 and amount < prev_m[0]["amount"] * 2:
            alerts.append(f"{merchant} subiu de {prev_m[0]['amount']:.2f} para {amount:.2f} {currency} "
                          f"em relação à cobrança de {prev_m[0]['date']}")
    # mês atual da categoria x média dos 3 meses anteriores
    month_start = d.replace(day=1)
    cur_total = sum(x["amount"] for x in store.query_expenses(month_start.isoformat(), d.isoformat(), category)
                    if x["currency"] == currency)
    prev = []
    ms = month_start
    for _ in range(3):
        pe = ms - timedelta(days=1)
        ms = pe.replace(day=1)
        tot = sum(x["amount"] for x in store.query_expenses(ms.isoformat(), pe.isoformat(), category)
                  if x["currency"] == currency)
        if tot:
            prev.append(tot)
    if prev:
        avg = sum(prev) / len(prev)
        if cur_total > avg * 1.3:
            alerts.append(f"{category} no mês já está em {cur_total:.2f} {currency}, acima da média "
                          f"de {avg:.2f} {currency} dos meses anteriores")
    return alerts


def period_compare(date_from: str, date_to: str, category=None, business=None) -> dict:
    a, b = date.fromisoformat(date_from), date.fromisoformat(date_to)
    out: dict = {}
    # mesmo intervalo no mês anterior
    def back(x: date) -> date:
        y, m = (x.year - 1, 12) if x.month == 1 else (x.year, x.month - 1)
        return date(y, m, min(x.day, _cal.monthrange(y, m)[1]))
    pa, pb = back(a), back(b)
    prev: dict = {}
    for r in store.query_expenses(pa.isoformat(), pb.isoformat(), category, business):
        prev[r["currency"]] = round(prev.get(r["currency"], 0) + r["amount"], 2)
    out["previous_period"] = {"from": pa.isoformat(), "to": pb.isoformat(), "totals_by_currency": prev}
    today = _now().date()
    if a.day == 1 and a.year == today.year and a.month == today.month and b >= today:
        days_in = _cal.monthrange(a.year, a.month)[1]
        cur: dict = {}
        for r in store.query_expenses(a.isoformat(), today.isoformat(), category, business):
            cur[r["currency"]] = cur.get(r["currency"], 0) + r["amount"]
        out["month_projection"] = {c: round(v / today.day * days_in, 2) for c, v in cur.items()}
    return out


# ---------- Bom dia ----------
def briefing_data(day: str | None = None) -> dict:
    from . import tools

    d = date.fromisoformat(day) if day else _now().date()
    data: dict = {"date": d.isoformat()}
    try:
        data["events"] = tools._list_events(f"{d}T00:00:00", f"{d}T23:59:59")["events"]
    except Exception as e:
        data["events_error"] = str(e)
    open_tasks = list_tasks("aberta")["tasks"]
    data["tasks_overdue"] = [t for t in open_tasks if t["due"] and t["due"] < d.isoformat()]
    data["tasks_today"] = [t for t in open_tasks if t["due"] == d.isoformat()]
    data["tasks_open_count"] = len(open_tasks)
    data["bills_next_7_days"] = [b for b in list_bills()["bills"]
                                 if b["next_due"] <= (d + timedelta(days=7)).isoformat()]
    data["documents_expiring_45_days"] = [
        {"title": r["title"], "expires_on": r["expires_on"]}
        for r in store.select("SELECT * FROM documents WHERE deleted=0 AND expires_on IS NOT NULL "
                              "AND expires_on>=? AND expires_on<=? ORDER BY expires_on",
                              (d.isoformat(), (d + timedelta(days=45)).isoformat()))]
    try:
        data["important_unread_emails"] = tools._search_emails("is:unread is:important newer_than:2d", 5)["emails"]
    except Exception:
        data["important_unread_emails"] = []
    month: dict = {}
    for r in store.query_expenses(d.replace(day=1).isoformat(), d.isoformat()):
        month[r["currency"]] = round(month.get(r["currency"], 0) + r["amount"], 2)
    data["spent_this_month"] = month
    return data


def daily_briefing(day=None):
    return briefing_data(day)


def briefing_text(data: dict) -> str:
    dias = ["segunda", "terça", "quarta", "quinta", "sexta", "sábado", "domingo"]
    d = date.fromisoformat(data["date"])
    hour = _now().hour
    hi = "Bom dia" if hour < 12 else "Boa tarde" if hour < 18 else "Boa noite"
    lines = [f"{hi}, {store.user_name()}. {dias[d.weekday()].capitalize()}, {d:%d/%m}."]
    ev = data.get("events") or []
    if ev:
        lines.append("\nAgenda:")
        for e in ev:
            t = (e.get("start") or "")[11:16] or "dia todo"
            lines.append(f"• {t} {e.get('title') or ''}" + (f" ({e['location'].split(',')[0]})" if e.get("location") else ""))
    else:
        lines.append("\nAgenda livre hoje.")
    if data["tasks_overdue"] or data["tasks_today"]:
        lines.append("\nTarefas:")
        lines += [f"• atrasada: {t['title']}" for t in data["tasks_overdue"]]
        lines += [f"• hoje: {t['title']}" for t in data["tasks_today"]]
    if data["bills_next_7_days"]:
        lines.append("\nContas nos próximos dias:")
        for b in data["bills_next_7_days"]:
            val = f" {b['amount']:.2f} {b['currency']}" if b.get("amount") else ""
            lines.append(f"• {b['name']}{val}, vence {date.fromisoformat(b['next_due']):%d/%m}")
    if data["documents_expiring_45_days"]:
        lines.append("\nDocumentos vencendo:")
        lines += [f"• {x['title']} em {date.fromisoformat(x['expires_on']):%d/%m}" for x in data["documents_expiring_45_days"]]
    em = data.get("important_unread_emails") or []
    if em:
        lines.append(f"\nE-mails importantes não lidos ({len(em)}):")
        lines += [f"• {e['from'].split('<')[0].strip()}: {e['subject']}" for e in em[:3]]
    if data["spent_this_month"]:
        lines.append("\nGastos do mês: " + ", ".join(f"{v:.2f} {c}" for c, v in data["spent_this_month"].items()))
    return "\n".join(lines)


# ---------- Resumo da semana ----------
def weekly_review_data() -> dict:
    from . import tools

    today = _now().date()
    start, prev_start = today - timedelta(days=7), today - timedelta(days=14)

    def spent(a: date, b: date) -> dict:
        out: dict = {}
        for r in store.query_expenses(a.isoformat(), b.isoformat()):
            out[r["currency"]] = round(out.get(r["currency"], 0) + r["amount"], 2)
        return out

    acts = store.select("SELECT kind, count(*) n FROM activities WHERE created_at>=? AND status NOT IN "
                        "('desfeito','cancelado') GROUP BY kind", (start.isoformat(),))
    done = store.select("SELECT title FROM tasks WHERE status='feita' AND done_at>=?", (start.isoformat(),))
    data = {"from": start.isoformat(), "to": today.isoformat(),
            "actions": {r["kind"]: r["n"] for r in acts},
            "tasks_done": [r["title"] for r in done],
            "tasks_open": len(list_tasks("aberta")["tasks"]),
            "spent_last_7_days": spent(start + timedelta(days=1), today),
            "spent_previous_7_days": spent(prev_start + timedelta(days=1), start),
            "bills_next_7_days": [b for b in list_bills()["bills"]
                                  if b["next_due"] <= (today + timedelta(days=7)).isoformat()]}
    try:
        data["next_7_days_events"] = tools._list_events(f"{today + timedelta(days=1)}T00:00:00",
                                                        f"{today + timedelta(days=7)}T23:59:59")["events"]
    except Exception:
        data["next_7_days_events"] = []
    return data


def weekly_review():
    return weekly_review_data()


def weekly_text(d: dict) -> str:
    total = sum(d["actions"].values())
    lines = [f"Sua semana com o Fidus ({date.fromisoformat(d['from']):%d/%m} a {date.fromisoformat(d['to']):%d/%m})"]
    lines.append(f"\n{total} coisas resolvidas por você" + (f", {len(d['tasks_done'])} tarefas concluídas." if d["tasks_done"] else "."))
    if d["spent_last_7_days"]:
        parts = []
        for c, v in d["spent_last_7_days"].items():
            p = d["spent_previous_7_days"].get(c)
            delta = f" ({'+' if v >= p else ''}{(v - p) / p * 100:.0f}% x semana anterior)" if p else ""
            parts.append(f"{v:.2f} {c}{delta}")
        lines.append("Gastos: " + ", ".join(parts))
    ev = d.get("next_7_days_events") or []
    if ev:
        lines.append(f"\nPróximos 7 dias: {len(ev)} compromissos")
        dias = ["seg", "ter", "qua", "qui", "sex", "sáb", "dom"]
        for e in ev[:5]:
            st = e.get("start") or ""
            try:
                dt = datetime.fromisoformat(st)
                when = f"{dias[dt.weekday()]} {dt:%d/%m}" + (f" {dt:%H:%M}" if len(st) > 10 else "")
            except ValueError:
                when = st
            lines.append(f"• {when} {e.get('title') or ''}")
    if d["bills_next_7_days"]:
        lines.append("Contas: " + ", ".join(f"{b['name']} ({date.fromisoformat(b['next_due']):%d/%m})" for b in d["bills_next_7_days"]))
    if d["tasks_open"]:
        lines.append(f"\n{d['tasks_open']} tarefas abertas na aba Tarefas.")
    return "\n".join(lines)


# ---------- Exportação para o contador ----------
def export_for_accountant(date_from, date_to, business=None):
    import os
    import zipfile

    from openpyxl import Workbook
    from openpyxl.styles import Font

    try:
        df, dt = date.fromisoformat(date_from), date.fromisoformat(date_to)
    except (TypeError, ValueError):
        return {"error": "datas precisam ser AAAA-MM-DD"}
    if business:
        b = store.match_business(business)
        if not b:  # carteira removida da lista, mas com gastos antigos guardados: também exporta
            old = store.select("SELECT business FROM expenses WHERE deleted=0 AND lower(business)=lower(?) LIMIT 1",
                               (business.strip(),))
            if not old:
                return {"error": f"carteira desconhecida: {business}. Carteiras: {', '.join(store.businesses())}"}
            b = old[0]["business"]
        business = b
    date_from, date_to = df.isoformat(), dt.isoformat()
    rows = sorted(store.query_expenses(date_from, date_to, business=business, limit=10000), key=lambda r: (r["date"], r["id"]))
    if not rows:
        return {"error": "nenhum gasto nesse período"}
    wb = Workbook()
    ws = wb.active
    ws.title = "Gastos"
    head = ["Data", "Empresa", "Categoria", "Estabelecimento", "Valor", "Moeda", "IVA/VAT", "Observação", "Recibo"]
    ws.append(head)
    for c in ws[1]:
        c.font = Font(bold=True)
    for r in rows:
        rec = f"recibos/{r['id']}{os.path.splitext(r['receipt_path'])[1]}" if r.get("receipt_path") and os.path.exists(r["receipt_path"]) else ""
        ws.append([r["date"], r["business"], r["category"], r["merchant"] or "", r["amount"], r["currency"],
                   r["vat"] if r["vat"] is not None else "", r["note"] or "", rec])
    for col, w in zip("ABCDEFGHI", (12, 18, 20, 24, 12, 8, 10, 30, 22)):
        ws.column_dimensions[col].width = w
    sm = wb.create_sheet("Resumo")
    sm.append(["Empresa", "Categoria", "Moeda", "Total", "IVA"])
    for c in sm[1]:
        c.font = Font(bold=True)
    agg: dict = {}
    for r in rows:
        k = (r["business"], r["category"], r["currency"])
        t = agg.setdefault(k, [0.0, 0.0])
        t[0] += r["amount"]
        t[1] += r["vat"] or 0
    for (b, c, cur), (tot, vat) in sorted(agg.items()):
        sm.append([b, c, cur, round(tot, 2), round(vat, 2)])
    folder = store.files_dir("exports")
    import re
    tag = re.sub(r"[^\w-]", "", (business or "todas").replace(" ", "-")) or "empresa"
    base = f"fidus-{tag}-{date_from}-a-{date_to}"
    xlsx = os.path.join(folder, base + ".xlsx")
    wb.save(xlsx)
    zpath = os.path.join(folder, base + ".zip")
    n_rec = 0
    with zipfile.ZipFile(zpath, "w", zipfile.ZIP_DEFLATED) as z:
        z.write(xlsx, base + ".xlsx")
        for r in rows:
            if r.get("receipt_path") and os.path.exists(r["receipt_path"]):
                z.write(r["receipt_path"], f"recibos/{r['id']}{os.path.splitext(r['receipt_path'])[1]}")
                n_rec += 1
    title = f"Pacote do contador · {business or 'todas as empresas'} · {date.fromisoformat(date_from):%d/%m/%y} a {date.fromisoformat(date_to):%d/%m/%y}"
    did = store.insert("documents", title=title, description=f"{len(rows)} gastos, {n_rec} recibos. Exportação.",
                       tags="exportação,contador", path=zpath, media_type="application/zip")
    return {"ok": True, "document_id": did, "title": title, "expenses": len(rows), "receipts": n_rec}


# ---------- Assinaturas ----------
def list_subscriptions():
    since = (_now().date() - timedelta(days=200)).isoformat()
    rows = store.select("SELECT * FROM expenses WHERE deleted=0 AND merchant IS NOT NULL AND merchant<>'' "
                        "AND date>=? ORDER BY date", (since,))
    groups: dict = {}
    for r in rows:
        groups.setdefault((r["merchant"].strip().lower(), r["currency"]), []).append(r)
    subs = []
    for (_m, cur), items in groups.items():
        months = {i["date"][:7] for i in items}
        if len(months) < 2:
            continue
        amounts = [i["amount"] for i in items]
        first, last = amounts[0], amounts[-1]
        if max(amounts) > min(amounts) * 1.5:  # valores muito diferentes: compras avulsas, não assinatura
            continue
        subs.append({"merchant": items[-1]["merchant"], "currency": cur, "current_amount": last,
                     "months_seen": len(months), "last_date": items[-1]["date"], "category": items[-1]["category"],
                     "price_change": round(last - first, 2) if abs(last - first) > 0.009 else 0})
    monthly: dict = {}
    for x in subs:
        monthly[x["currency"]] = round(monthly.get(x["currency"], 0) + x["current_amount"], 2)
    return {"subscriptions": sorted(subs, key=lambda x: -x["current_amount"]), "monthly_total_by_currency": monthly}


# ---------- Contador para a aba Atividade ----------
def month_stats() -> dict:
    start = _now().date().replace(day=1).isoformat()
    rows = store.select("SELECT kind, count(*) n FROM activities WHERE created_at>=? AND status NOT IN "
                        "('desfeito','cancelado') GROUP BY kind", (start,))
    by = {r["kind"]: r["n"] for r in rows}
    total = sum(by.values())
    # estimativa conservadora: ~3 min por ação que você não precisou fazer
    return {"actions_this_month": total, "by_kind": by, "minutes_saved_estimate": total * 3}


DISPATCH = {
    "weekly_review": lambda: weekly_review(),
    "export_for_accountant": lambda **a: export_for_accountant(**a),
    "list_subscriptions": lambda: list_subscriptions(),
    "create_reminder": create_reminder,
    "add_meet_link": add_meet_link,
    "prepare_event_invite": prepare_event_invite,
    "add_task": add_task,
    "list_tasks": list_tasks,
    "complete_task": complete_task,
    "search_documents": search_documents,
    "add_bill": add_bill,
    "list_bills": list_bills,
    "delete_bill": delete_bill,
    "daily_briefing": daily_briefing,
}
